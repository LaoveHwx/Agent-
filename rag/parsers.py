"""上传文件解析：把二进制文件抽成纯文本，供 RAG 切分入库。

支持的格式：
- .pdf  -> pypdf 逐页抽取
- .docx -> python-docx 段落拼接
- .doc  -> 老 .doc 二进制格式 python-docx 无法直接解析，先尝试按 docx 读，
           失败则抛出明确错误，提示另存为 .docx
- .txt/.md/.markdown/.csv/.json/.log -> 多编码文本解码
"""
from __future__ import annotations

import io
from typing import Callable

from utils.logger import setup_logger


logger = setup_logger(__name__)


# 纯文本类后缀：直接按编码解码
_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".json", ".log"}

# 二进制类后缀：需要专用解析器
_PARSERS: dict[str, Callable[[bytes], str]] = {
    ".pdf": None,   # 延迟填充，见下
    ".docx": None,
    ".doc": None,
}


def _decode_text(data: bytes) -> str:
    """按常见编码依次尝试解码，兜底用 utf-8 容错。"""
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_pdf(data: bytes) -> str:
    """逐页抽取 PDF 文本并用换行拼接。"""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception:
            logger.exception("pdf page extract failed, page=%s", index)
            text = ""
        if text:
            parts.append(text)
    return "\n".join(parts)


def _parse_docx(data: bytes) -> str:
    """抽取 docx 段落与表格单元格文本并用换行拼接。"""
    from docx import Document

    document = Document(io.BytesIO(data))
    # 段落 + 表格单元格一起抽，避免漏掉表格内容
    parts: list[str] = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def _parse_doc(data: bytes) -> str:
    """老 .doc 二进制格式无纯 Python 解析器，先按 docx 试，失败给明确提示。"""
    try:
        return _parse_docx(data)
    except Exception as exc:
        raise ValueError("legacy .doc 格式暂不支持，请另存为 .docx 后重试") from exc


_PARSERS[".pdf"] = _parse_pdf
_PARSERS[".docx"] = _parse_docx
_PARSERS[".doc"] = _parse_doc


SUPPORTED_SUFFIXES = _TEXT_SUFFIXES | set(_PARSERS.keys())


def file_suffix(filename: str) -> str:
    """取文件名后缀（含点并小写），无后缀返回空串。"""
    dot_index = filename.rfind(".")
    return filename[dot_index:].lower() if dot_index >= 0 else ""


def is_supported(filename: str) -> bool:
    """判断文件名后缀是否在支持解析的集合内。"""
    return file_suffix(filename) in SUPPORTED_SUFFIXES


def parse_upload(filename: str, data: bytes) -> str:
    """根据文件名后缀选择解析器，把 bytes 抽成纯文本。"""
    suffix = file_suffix(filename)
    parser = _PARSERS.get(suffix)
    if parser is not None:
        return parser(data)
    # 其余按文本解码（含 _TEXT_SUFFIXES 以及无后缀的兜底）
    return _decode_text(data)
